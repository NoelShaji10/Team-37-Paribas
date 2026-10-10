"""
Reads one uploaded file, counts its transaction rows and measures how long that took.

This module knows nothing about the web: give it a file path, get back a plain dict (a "record").
It NEVER raises for a bad file - problems become a record with status="failure" and a message.
That makes it easy to test and easy to reuse later (e.g. by the prediction code in Phase 2).

What is timed: parsing + counting only (not the network upload).
"""
from __future__ import annotations

import csv
import json
import time
import uuid
from datetime import datetime
from pathlib import Path

SUPPORTED_FORMATS = {".csv": "CSV", ".json": "JSON", ".xlsx": "XLSX"}
JSON_LIST_KEYS = ("transactions", "records", "data", "items", "rows", "results")
SAMPLE_PREFIXES = ("sample_", "synthetic_")      # files named like this are labelled as sample data
MAX_ERROR_LEN = 300


class ProcessingError(Exception):
    """A problem with the file's contents. The message is safe to show to the user."""


# ------------------------------------------------------------------ row counters (one per format)
def count_csv_rows(path: Path) -> int:
    """Streams the file row by row (constant memory, fine for very large CSVs).
    First non-blank row = header; every other non-blank row = one transaction."""
    with open(path, "rb") as f:
        if b"\x00" in f.read(4096):
            raise ProcessingError("File contains binary data and does not look like a CSV text file")
    rows, header_seen = 0, False
    try:
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            for rec in csv.reader(f):
                if not rec or all(not cell.strip() for cell in rec):
                    continue                                   # blank line
                if not header_seen:
                    header_seen = True
                    continue
                rows += 1
    except UnicodeDecodeError:
        raise ProcessingError("File is not valid UTF-8 text") from None
    except csv.Error as e:
        raise ProcessingError(f"Malformed CSV: {e}") from None
    return rows


def _find_records(data):
    """JSON can be a list of transactions, or an object that holds that list."""
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for key in JSON_LIST_KEYS:
            if isinstance(data.get(key), list):
                return data[key]
        lists = [v for v in data.values() if isinstance(v, list)]
        if len(lists) == 1:
            return lists[0]
    raise ProcessingError("Could not find a list of transactions in the JSON "
                          "(expected a list, or an object with a 'transactions' list)")


def count_json_rows(path: Path) -> int:
    """NOTE: json.load reads the whole file into memory (the standard library cannot stream JSON).
    Fine for files up to a few hundred MB; for bigger ones use CSV."""
    try:
        with open(path, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
    except UnicodeDecodeError:
        raise ProcessingError("File is not valid UTF-8 text") from None
    except json.JSONDecodeError as e:
        raise ProcessingError(f"Invalid JSON: {e.msg} (line {e.lineno}, column {e.colno})") from None
    except RecursionError:
        raise ProcessingError("JSON is nested too deeply") from None
    except MemoryError:
        raise ProcessingError("JSON file is too large to load; please use CSV for very large files") from None
    return len(_find_records(data))


def count_xlsx_rows(path: Path) -> int:
    """Reads the first (active) sheet in streaming 'read-only' mode. First non-empty row = header."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ProcessingError("Excel support needs the openpyxl package (pip install openpyxl)") from None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise ProcessingError("Not a valid .xlsx file (it may be corrupt or a different format)") from None
    rows, header_seen = 0, False
    try:
        ws = wb.active
        if ws is None:
            raise ProcessingError("Workbook has no readable sheet")
        for rec in ws.iter_rows(values_only=True):
            if all(c is None or (isinstance(c, str) and not c.strip()) for c in rec):
                continue
            if not header_seen:
                header_seen = True
                continue
            rows += 1
    except ProcessingError:
        raise
    except Exception as e:
        raise ProcessingError(f"Could not read the Excel sheet: {e}") from None
    finally:
        wb.close()
    return rows


COUNTERS = {"CSV": count_csv_rows, "JSON": count_json_rows, "XLSX": count_xlsx_rows}


# ------------------------------------------------------------------ records
def _fmt_time(ts: float) -> str:
    return datetime.fromtimestamp(ts).isoformat(sep=" ", timespec="milliseconds")


def is_sample(file_name: str) -> bool:
    return file_name.lower().startswith(SAMPLE_PREFIXES)


def _record(file_name, file_format, size_bytes, rows, start_ts, end_ts, duration, status, error):
    return {
        "file_id": uuid.uuid4().hex[:8],
        "file_name": file_name,
        "file_format": file_format,
        "size_bytes": size_bytes,
        "size_kb": round(size_bytes / 1024, 2),
        "rows": rows,
        "start_time": _fmt_time(start_ts),
        "end_time": _fmt_time(end_ts),
        "duration_sec": round(duration, 4),
        "status": status,
        "error_message": error[:MAX_ERROR_LEN],
        "data_label": "sample (synthetic)" if is_sample(file_name) else "user upload",
    }


def failure_record(file_name: str, size_bytes: int, message: str) -> dict:
    """A failure that happened before processing could start (e.g. upload too large)."""
    now = time.time()
    name = Path(file_name).name or "unnamed"
    fmt = SUPPORTED_FORMATS.get(Path(name).suffix.lower()) or (Path(name).suffix.lstrip(".").upper() or "UNKNOWN")
    return _record(name, fmt, size_bytes, 0, now, now, 0.0, "failure", message)


def process_file(path: Path, original_name: str) -> dict:
    """Parse the file at `path`, count transaction rows, time it. Never raises."""
    name = Path(original_name).name or "unnamed"
    ext = Path(name).suffix.lower()
    fmt = SUPPORTED_FORMATS.get(ext)
    shown_format = fmt or (ext.lstrip(".").upper() or "UNKNOWN")
    size = Path(path).stat().st_size

    start_ts, t0 = time.time(), time.perf_counter()
    rows, status, error = 0, "success", ""
    try:
        if fmt is None:
            raise ProcessingError(f"Unsupported file type '{ext or '(none)'}'. Allowed: .csv, .json, .xlsx")
        if size == 0:
            raise ProcessingError("File is empty (0 bytes)")
        rows = COUNTERS[fmt](Path(path))
        if rows == 0:
            raise ProcessingError("File contains no transaction rows")
    except ProcessingError as e:
        status, error = "failure", str(e)
    except Exception as e:                                   # anything unexpected: record it, do not crash
        status, error = "failure", f"Unexpected error ({e.__class__.__name__}): {e}"
    duration = time.perf_counter() - t0
    return _record(name, shown_format, size, rows, start_ts, time.time(), duration, status, error)
